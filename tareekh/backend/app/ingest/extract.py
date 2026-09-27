"""Step 1 of ingest: get plain text out of whatever was uploaded."""
import io
import mimetypes
from pathlib import Path

from .. import llm

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".heic"}
TEXT_EXT = {".txt", ".md"}


def kind_of(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in IMAGE_EXT:
        return "image"
    if ext in TEXT_EXT:
        return "text"
    if ext == ".docx":
        return "docx"
    if ext == ".pdf":
        return "pdf"
    return "unknown"


def extract(path: Path) -> tuple[str, str]:
    """Return (text, method). method is recorded so the UI can say 'OCR' vs 'text layer'."""
    kind = kind_of(path.name)
    if kind == "text":
        return path.read_text(encoding="utf-8", errors="replace"), "text"
    if kind == "docx":
        from docx import Document
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip()), "docx"
    if kind == "pdf":
        return _pdf(path)
    if kind == "image":
        mime = mimetypes.guess_type(path.name)[0] or "image/jpeg"
        return llm.ocr_image(path.read_bytes(), mime), "ocr"
    raise ValueError(f"Unsupported file type: {path.name}")


def _pdf(path: Path) -> tuple[str, str]:
    from pypdf import PdfReader
    text = "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages).strip()
    if len(text) > 40:
        return text, "pdf-text"
    # scanned PDF: render pages and OCR them
    import pypdfium2 as pdfium
    pages = []
    for page in pdfium.PdfDocument(str(path)):
        buf = io.BytesIO()
        page.render(scale=2).to_pil().save(buf, "JPEG", quality=85)
        pages.append(llm.ocr_image(buf.getvalue()))
    return "\n\n".join(pages), "pdf-ocr"
